#!/usr/bin/env python3
"""
Read BOM data from Altium export XLSX files.

Two BOM files are expected in the export folder:
  - Bill of Materials*.xlsx  — main BOM with value, MPN, manufacturer
  - *Tranfer.xlsx or *Transfer.xlsx — Highstage transfer BOM with internal part numbers

Returns a component_map matching the structure produced by
pdf_annotation_extractor.build_component_map(), keyed by refdes.
"""

from pathlib import Path


def _find_file(folder: Path, pattern: str) -> Path | None:
    matches = list(folder.glob(pattern))
    return matches[0] if matches else None


def _split_designators(raw: str) -> list[str]:
    """Split a comma-separated designator string into individual refdes strings."""
    if not raw:
        return []
    return [r.strip() for r in str(raw).split(',') if r.strip()]


def _read_main_bom(bom_path: Path, component_map: dict) -> None:
    """Populate component_map from the main Bill of Materials XLSX."""
    import openpyxl
    wb = openpyxl.load_workbook(bom_path, read_only=True, data_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    # Find header row — the first row containing 'Designator'
    header_idx = None
    for i, row in enumerate(rows):
        row_strs = [str(c).strip() if c is not None else '' for c in row]
        if 'Designator' in row_strs:
            header_idx = i
            break

    if header_idx is None:
        return

    headers = [str(c).strip() if c is not None else '' for c in rows[header_idx]]

    col = {h: i for i, h in enumerate(headers) if h}
    des_col  = col.get('Designator')
    val_col  = col.get('Name')
    mpn_col  = col.get('source_manufacturer_1_pn')
    mfg_col  = col.get('source_manufacturer_1')
    desc_col = col.get('Description')

    if des_col is None:
        return

    for row in rows[header_idx + 1:]:
        des_raw = row[des_col] if des_col < len(row) else None
        if des_raw is None:
            continue

        value   = str(row[val_col]).strip()  if val_col  is not None and row[val_col]  is not None else ''
        mpn     = str(row[mpn_col]).strip()  if mpn_col  is not None and row[mpn_col]  is not None else ''
        mfg     = str(row[mfg_col]).strip()  if mfg_col  is not None and row[mfg_col]  is not None else ''
        desc    = str(row[desc_col]).strip() if desc_col is not None and row[desc_col] is not None else ''

        entry: dict = {}
        if value:
            entry['value'] = value
        if mpn:
            entry['mpn1'] = mpn
        if mfg:
            entry['mfg1'] = mfg
        if desc:
            entry['description'] = desc

        if not entry:
            continue

        for ref in _split_designators(des_raw):
            if ref not in component_map:
                component_map[ref] = {}
            component_map[ref].update(entry)


def _read_transfer_bom(bom_path: Path, component_map: dict) -> None:
    """Merge Highstage part numbers from the transfer BOM into component_map."""
    import openpyxl
    wb = openpyxl.load_workbook(bom_path, read_only=True, data_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    # Find header row — the first row containing 'Designator'
    header_idx = None
    for i, row in enumerate(rows):
        row_strs = [str(c).strip() if c is not None else '' for c in row]
        if 'Designator' in row_strs:
            header_idx = i
            break

    if header_idx is None:
        return

    headers = [str(c).strip() if c is not None else '' for c in rows[header_idx]]
    col = {h: i for i, h in enumerate(headers) if h}

    des_col  = col.get('Designator')
    hs_col   = col.get('Highstage replacement')
    mpn_col  = col.get('source_manufacturer_1_pn')
    mfg_col  = col.get('source_manufacturer_1')
    val_col  = col.get('Comment')
    desc_col = col.get('Description')

    if des_col is None:
        return

    for row in rows[header_idx + 1:]:
        des_raw = row[des_col] if des_col < len(row) else None
        if des_raw is None:
            continue

        hs_pn = str(row[hs_col]).strip()   if hs_col  is not None and row[hs_col]  is not None else ''
        mpn   = str(row[mpn_col]).strip()  if mpn_col is not None and row[mpn_col] is not None else ''
        mfg   = str(row[mfg_col]).strip()  if mfg_col is not None and row[mfg_col] is not None else ''
        value = str(row[val_col]).strip()  if val_col is not None and row[val_col]  is not None else ''
        desc  = str(row[desc_col]).strip() if desc_col is not None and row[desc_col] is not None else ''

        for ref in _split_designators(des_raw):
            entry = component_map.setdefault(ref, {})
            if hs_pn and not entry.get('highstage_pn'):
                entry['highstage_pn'] = hs_pn
            if mpn and not entry.get('mpn1'):
                entry['mpn1'] = mpn
            if mfg and not entry.get('mfg1'):
                entry['mfg1'] = mfg
            if value and not entry.get('value'):
                entry['value'] = value
            if desc and not entry.get('description'):
                entry['description'] = desc


def read_altium_bom(export_folder: Path) -> dict[str, dict]:
    """Read BOM data from Altium export folder. Returns component_map keyed by refdes."""
    component_map: dict[str, dict] = {}

    main_bom = _find_file(export_folder, 'Bill of Materials*.xlsx')
    if main_bom:
        _read_main_bom(main_bom, component_map)

    transfer_bom = (_find_file(export_folder, '*Tranfer.xlsx')
                    or _find_file(export_folder, '*Transfer.xlsx'))
    if transfer_bom:
        _read_transfer_bom(transfer_bom, component_map)

    return component_map
